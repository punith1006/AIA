"""
Generate a formatted Word document from the security audit walkthrough.
"""

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

def set_cell_shading(cell, color):
    """Set cell background color"""
    shading_elm = OxmlElement('w:shd')
    shading_elm.set(qn('w:fill'), color)
    cell._tc.get_or_add_tcPr().append(shading_elm)

def add_heading_with_style(doc, text, level):
    """Add a heading with custom styling"""
    heading = doc.add_heading(text, level)
    if level == 1:
        heading.runs[0].font.color.rgb = RGBColor(0, 51, 102)
    return heading

def add_table(doc, headers, rows, header_color="1F4E79"):
    """Add a formatted table"""
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = 'Table Grid'
    
    # Header row
    header_cells = table.rows[0].cells
    for i, header in enumerate(headers):
        header_cells[i].text = header
        header_cells[i].paragraphs[0].runs[0].bold = True
        header_cells[i].paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
        set_cell_shading(header_cells[i], header_color)
    
    # Data rows
    for row_data in rows:
        row_cells = table.add_row().cells
        for i, cell_data in enumerate(row_data):
            row_cells[i].text = str(cell_data)
    
    doc.add_paragraph()  # Add spacing after table
    return table

def create_docx():
    doc = Document()
    
    # Title
    title = doc.add_heading('Deep Dive Analysis: Server.py & Agent Architecture', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # Subtitle
    subtitle = doc.add_paragraph('Security Audit & Architectural Review')
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.runs[0].font.size = Pt(14)
    subtitle.runs[0].font.color.rgb = RGBColor(128, 128, 128)
    
    doc.add_paragraph()
    
    # ========== PROJECT OVERVIEW ==========
    add_heading_with_style(doc, '1. Project Overview', 1)
    doc.add_paragraph(
        'This is a Python-based Google ADK (Agent Development Kit) agent system for automated '
        'business intelligence research. The system provides a unified FastAPI server that '
        'orchestrates AI-powered research agents for organizational and market intelligence.'
    )
    
    add_heading_with_style(doc, 'Current Project Structure', 2)
    structure = doc.add_paragraph()
    structure.add_run('''aia_mk_org/
├── server.py                    # Main FastAPI server (Port 8500)
├── .env                         # Environment variables
├── req_aia.txt                  # Python dependencies
├── logs/                        # Server logs directory
├── org_research/                # Organizational Research Agent
│   ├── agent.py                 # Root agent definition
│   ├── config.py                # Model configuration
│   ├── sub_agents/              # Client research pipeline
│   └── tools/mongoupload.py     # MongoDB operations + HTTP calls
└── market_stream_uat/           # Market Intelligence Agent
    ├── agent.py                 # Root agent (565 lines)
    ├── config.py                # Model configuration
    ├── .env                     # Local environment variables
    ├── sub_agents/              # 5 sub-agent modules
    └── tools/mongoupload.py     # MongoDB operations + HTTP calls''').font.name = 'Consolas'
    
    # ========== EXPOSED ENDPOINTS ==========
    add_heading_with_style(doc, '2. Exposed Endpoints (Inbound - Port 8500)', 1)
    
    add_heading_with_style(doc, 'REST Endpoints', 2)
    add_table(doc, 
        ['Method', 'Path', 'Source', 'Auth', 'Purpose'],
        [
            ['GET', '/', 'server.py:375', '❌ None', 'Server info JSON with endpoints'],
            ['GET', '/health', 'server.py:391', '❌ None', 'Health check (sessions, tasks)'],
            ['GET', '/sessions', 'server.py:402', '❌ None', 'Debug: View active sessions'],
        ]
    )
    
    add_heading_with_style(doc, 'WebSocket Endpoints', 2)
    add_table(doc,
        ['Protocol', 'Path', 'Source', 'Auth', 'Target Agent'],
        [
            ['WS', '/ws/org-research', 'server.py:356', '❌ None', 'org_research agent'],
            ['WS', '/ws/market-research', 'server.py:361', '❌ None', 'market_research agent'],
            ['WS', '/run-live-agent', 'server.py:367', '❌ None', 'org_research (legacy)'],
        ]
    )
    
    add_heading_with_style(doc, 'ADK Auto-Mounted Endpoints', 2)
    add_table(doc,
        ['Path Pattern', 'Purpose'],
        [
            ['/dev-ui', 'ADK Developer Web UI'],
            ['/run', 'ADK agent execution API'],
            ['/events', 'SSE event stream'],
            ['/{agent}/run', 'Agent-specific execution'],
        ]
    )
    
    # ========== INTERNAL API CALLS ==========
    add_heading_with_style(doc, '3. Internal API Calls (Outbound to localhost:4005)', 1)
    
    doc.add_paragraph(
        'These are outbound HTTP calls made by the Python agents to an external Node.js backend '
        'service running on port 4005. This service manages project status tracking, onboarding '
        'workflow notifications, and UI status updates.'
    )
    
    add_table(doc,
        ['Source File', 'Line', 'Method', 'Target URL', 'Payload', 'Trigger'],
        [
            ['org_research/tools/mongoupload.py', '54', 'POST', 
             'http://localhost:4005/onboard/update-org-report/{client_id}/', 
             'None', 'After org report saved'],
            ['market_stream_uat/tools/mongoupload.py', '73', 'PUT', 
             'http://localhost:4005/project/project-status-update/{project_id}/', 
             '{"sub_status": "..."}', 'After each report saved'],
            ['market_stream_uat/agent.py', '513', 'PUT', 
             'http://localhost:4005/project/project-status-update/{project_id}/', 
             '{"sub_status": "Completed"}', 'Pipeline finished'],
        ]
    )
    
    # ========== DATABASE CONNECTIONS ==========
    add_heading_with_style(doc, '4. Database Connections (MongoDB)', 1)
    
    add_heading_with_style(doc, 'Connection String', 2)
    conn = doc.add_paragraph()
    conn.add_run('MONGO_DB_CONNECTOR=mongodb://localhost:27017').font.name = 'Consolas'
    
    add_heading_with_style(doc, 'Database Operations', 2)
    add_table(doc,
        ['Source File', 'Function', 'Database', 'Collection', 'Operation'],
        [
            ['org_research/tools/mongoupload.py', 'create_blank_project()', 'sales_reports', 'org_reports', 'INSERT'],
            ['org_research/tools/mongoupload.py', 'update_project_report()', 'sales_reports', 'org_reports', 'UPDATE'],
            ['market_stream_uat/tools/mongoupload.py', 'create_blank_project()', 'sales_reports', 'projects', 'INSERT'],
            ['market_stream_uat/tools/mongoupload.py', 'update_project_report()', 'sales_reports', 'projects', 'UPDATE'],
        ]
    )
    
    add_heading_with_style(doc, 'Document Schemas', 2)
    
    doc.add_paragraph('org_reports Collection:', style='Intense Quote')
    schema1 = doc.add_paragraph()
    schema1.add_run('''{
  "client_id": "string",
  "client_org_research": {
    "html_report": "...",
    "raw_report": "..."
  }
}''').font.name = 'Consolas'
    
    doc.add_paragraph('projects Collection:', style='Intense Quote')
    schema2 = doc.add_paragraph()
    schema2.add_run('''{
  "project_id": "string",
  "client_org_research": "",
  "prospect_research": "",
  "market_segment": "",
  "target_org_research": "",
  "market_context": "",
  "client_org_research_html": "",
  "market_context_html": "",
  "prospect_research_html": "",
  "market_segment_html": "",
  "target_org_research_html": ""
}''').font.name = 'Consolas'
    
    # ========== EXTERNAL SERVICES ==========
    add_heading_with_style(doc, '5. External Service Connections', 1)
    
    add_table(doc,
        ['Service', 'Protocol', 'Endpoint', 'Auth Key', 'Purpose'],
        [
            ['Google Gemini API', 'HTTPS', 'generativelanguage.googleapis.com', 'GOOGLE_API_KEY', 'AI model inference + search'],
            ['OpenAI API', 'HTTPS', 'api.openai.com', 'OPENAI_API_KEY', 'gpt-5-mini, gpt-5-nano via LiteLlm'],
            ['Google Search Grounding', 'HTTPS', 'Google Search API', 'GOOGLE_API_KEY', 'Web research grounding'],
        ]
    )
    
    add_heading_with_style(doc, 'Agents Using google_search Tool', 2)
    add_table(doc,
        ['File', 'Usage Locations (Lines)'],
        [
            ['org_research/.../client_research_agent.py', '618, 898, 1063'],
            ['market_stream_uat/.../target_research.py', '306, 518, 640'],
            ['market_stream_uat/.../segment_agent.py', '220, 576, 685'],
            ['market_stream_uat/.../prospect_agent.py', '167, 352'],
            ['market_stream_uat/.../market_context_agent.py', '260, 551, 693'],
            ['market_stream_uat/.../client_research_agent.py', '573, 715, 873'],
        ]
    )
    
    # ========== ENVIRONMENT VARIABLES ==========
    add_heading_with_style(doc, '6. Environment Variables', 1)
    
    add_table(doc,
        ['File', 'Variable', 'Value/Description'],
        [
            ['.env (root)', 'GOOGLE_GENAI_USE_VERTEXAI', 'FALSE'],
            ['.env (root)', 'GOOGLE_API_KEY', 'AIzaSy... (exposed)'],
            ['.env (root)', 'MONGO_DB_CONNECTOR', 'mongodb://localhost:27017'],
            ['.env (root)', 'OPENAI_API_KEY', 'sk-proj-... (exposed)'],
            ['market_stream_uat/.env', 'GOOGLE_API_KEY', 'AIzaSy... (same key)'],
            ['market_stream_uat/.env', 'OPENAI_API_KEY', 'sk-proj-... (same key)'],
        ]
    )
    
    warning = doc.add_paragraph()
    warning.add_run('⚠️ SECURITY RISK: ').bold = True
    warning.add_run('API keys are stored in plaintext in .env files. Ensure these are NOT committed to version control.')
    
    # ========== FILE ANALYSIS ==========
    add_heading_with_style(doc, '7. File-by-File Analysis', 1)
    
    add_heading_with_style(doc, 'server.py (17,599 bytes, 447 lines)', 2)
    add_table(doc,
        ['Aspect', 'Details'],
        [
            ['Purpose', 'Main entry point - FastAPI unified server'],
            ['Port', '8500'],
            ['CORS', 'allow_origins=["*"] (permissive)'],
            ['Session Management', 'InMemorySessionService'],
            ['Stagnation Timeout', '300 seconds'],
            ['Logging', 'Console + logs/unified_server_YYYYMMDD.log'],
        ]
    )
    
    add_heading_with_style(doc, 'org_research/agent.py (8,037 bytes)', 2)
    add_table(doc,
        ['Aspect', 'Details'],
        [
            ['Root Agent', 'organizational_only_agent (SequentialAgent)'],
            ['Sub-agents', 'input_analyzer → project_creator → org_prompt_builder → organizational_intelligence_agent'],
            ['Tools', 'create_blank_project'],
            ['Callbacks', 'extract_client_id, store_organizational_report'],
        ]
    )
    
    add_heading_with_style(doc, 'market_stream_uat/agent.py (27,395 bytes, 565 lines)', 2)
    add_table(doc,
        ['Aspect', 'Details'],
        [
            ['Root Agent', 'simplified_intelligence_agent (SequentialAgent)'],
            ['Sub-agents', '14 agents in sequence (input_analyzer → harbinger)'],
            ['Callbacks', 'store_context_report, store_segmentation_report, store_target_report, store_prospect_report, harbinger_message'],
        ]
    )
    
    add_heading_with_style(doc, 'Model Configuration', 2)
    add_table(doc,
        ['Model', 'Value', 'Purpose'],
        [
            ['critic_model', 'LiteLlm("openai/gpt-5-mini")', 'Evaluation/quality checks'],
            ['worker_model', 'LiteLlm("openai/gpt-5-nano")', 'General operations'],
            ['search_model', 'Gemini("gemini-2.5-flash")', 'Web search grounding'],
            ['max_search_iterations', '2 (org_research), 0 (market_stream_uat)', 'Quality loop limit'],
        ]
    )
    
    # ========== SECURITY SUMMARY ==========
    add_heading_with_style(doc, '8. Security Findings Summary', 1)
    
    add_table(doc,
        ['Issue', 'Severity', 'Details'],
        [
            ['No Authentication', '🔴 Critical', 'All endpoints publicly accessible'],
            ['Permissive CORS', '🔴 Critical', 'allow_origins=["*"] allows any origin'],
            ['API Keys in .env', '🟠 High', 'Keys in plaintext (2 .env files)'],
            ['Debug Endpoints', '🟠 High', '/sessions exposes internal state'],
            ['No MongoDB Auth', '🟡 Medium', 'mongodb://localhost:27017 with no credentials'],
            ['HTTP to localhost:4005', '🟡 Medium', 'Unencrypted internal communication'],
            ['Binding 0.0.0.0', '🟡 Medium', 'Server listens on all interfaces'],
        ],
        header_color="8B0000"
    )
    
    # ========== COMPLETE SUMMARY ==========
    add_heading_with_style(doc, '9. Complete Connection Summary', 1)
    
    add_heading_with_style(doc, 'All Exposed Endpoints (Inbound)', 2)
    add_table(doc,
        ['Type', 'Method', 'Path', 'Port', 'Auth', 'Purpose'],
        [
            ['REST', 'GET', '/', '8500', '❌', 'Server info'],
            ['REST', 'GET', '/health', '8500', '❌', 'Health check'],
            ['REST', 'GET', '/sessions', '8500', '❌', 'Debug sessions'],
            ['WebSocket', 'WS', '/ws/org-research', '8500', '❌', 'Org research'],
            ['WebSocket', 'WS', '/ws/market-research', '8500', '❌', 'Market research'],
            ['WebSocket', 'WS', '/run-live-agent', '8500', '❌', 'Legacy endpoint'],
            ['ADK', '*', '/dev-ui', '8500', '❌', 'Developer UI'],
            ['ADK', '*', '/run', '8500', '❌', 'Agent execution'],
        ]
    )
    
    add_heading_with_style(doc, 'All Outbound Connections', 2)
    add_table(doc,
        ['Destination', 'Protocol', 'Port/Endpoint', 'Purpose'],
        [
            ['localhost:4005', 'HTTP', '/onboard/update-org-report/', 'Notify org report ready'],
            ['localhost:4005', 'HTTP', '/project/project-status-update/', 'Update project status'],
            ['localhost:27017', 'MongoDB', 'sales_reports database', 'Data persistence'],
            ['Google Gemini API', 'HTTPS', 'generativelanguage.googleapis.com', 'AI inference'],
            ['OpenAI API', 'HTTPS', 'api.openai.com', 'GPT models'],
            ['Google Search API', 'HTTPS', 'Google Search', 'Web grounding'],
        ]
    )
    
    # Footer
    doc.add_paragraph()
    footer = doc.add_paragraph('Document generated from security audit walkthrough.')
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.runs[0].font.color.rgb = RGBColor(128, 128, 128)
    
    # Save document
    output_path = r'c:\Users\Punith\Downloads\aia_mk_org\aia_mk_org\Security_Audit_Walkthrough.docx'
    doc.save(output_path)
    print(f"Document saved to: {output_path}")
    return output_path

if __name__ == "__main__":
    create_docx()
